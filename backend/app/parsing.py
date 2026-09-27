import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path

import docx
import pypdfium2 as pdfium
from docx.table import Table
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from pptx.shapes.group import GroupShape

from app.config import MAX_UNZIPPED_MB
from app.schemas import DocType, ErrorCode

MAX_LOCATION_CHARS = 120  # BACKEND_SCHEMA §4
DOC_TYPES: dict[str, DocType] = {".pdf": "pdf", ".docx": "docx", ".pptx": "pptx"}
ZIP_MAIN_PART = {"docx": "word/document.xml", "pptx": "ppt/presentation.xml"}


class UploadError(Exception):
    def __init__(self, code: ErrorCode):
        super().__init__(code)
        self.code = code


@dataclass(frozen=True)
class Section:
    location: str  # display label: "p. 12", "slide 4", "§ Normal forms"
    order: int     # page / slide / section number, 1-based
    text: str


@dataclass(frozen=True)
class Chunk:
    location: str
    order: int
    chunk_index: int  # 0-based position within the document
    text: str


def detect_type(path: Path | str, filename: str) -> DocType:
    ext = Path(filename).suffix.lower()
    if ext in (".doc", ".ppt"):
        raise UploadError("legacy_format")
    doc_type = DOC_TYPES.get(ext)
    if doc_type is None:
        raise UploadError("unsupported_type")

    if doc_type == "pdf":
        with open(path, "rb") as f:
            if f.read(5) != b"%PDF-":
                raise UploadError("bad_signature")
        return doc_type

    try:
        with zipfile.ZipFile(path) as z:
            entries = z.infolist()
    except (zipfile.BadZipFile, NotImplementedError):  # NotImplementedError: entry needs a ZIP version zipfile lacks
        raise UploadError("bad_signature") from None
    if ZIP_MAIN_PART[doc_type] not in {e.filename for e in entries}:
        raise UploadError("bad_signature")
    # zipfile never inflates an entry past its declared size, so capping the declared
    # total also caps what the parsers can decompress (zip-bomb guard).
    if sum(e.file_size for e in entries) > MAX_UNZIPPED_MB * 1024 * 1024:
        raise UploadError("bad_signature")
    return doc_type


def parse(path: Path | str, doc_type: DocType) -> list[Section]:
    return {"pdf": parse_pdf, "docx": parse_docx, "pptx": parse_pptx}[doc_type](path)


def parse_pdf(path: Path | str) -> list[Section]:
    pdf = pdfium.PdfDocument(path)
    try:
        sections = []
        for i in range(len(pdf)):
            page = pdf[i]
            textpage = page.get_textpage()
            text = textpage.get_text_range().replace("\r\n", "\n")
            # Close per page: a 500-page PDF would otherwise hold every page in memory on a 512 MB host.
            textpage.close()
            page.close()
            sections.append(Section(f"p. {i + 1}", i + 1, text))
        return sections
    finally:
        pdf.close()


def parse_docx(path: Path | str) -> list[Section]:
    sections: list[Section] = []
    heading, lines = "(start)", []

    def close_section() -> None:
        text = "\n".join(lines)
        if text.strip():
            location = f"§ {heading}"[:MAX_LOCATION_CHARS]
            sections.append(Section(location, len(sections) + 1, text))

    for block in docx.Document(str(path)).iter_inner_content():
        if isinstance(block, Table):
            lines.extend(_table_lines(_docx_cells(row) for row in block.rows))
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name if block.style else ""
        if style.startswith("Heading") or style == "Title":
            close_section()
            heading, lines = " ".join(text.split()), [text]
        else:
            lines.append(text)
    close_section()
    return sections


def parse_pptx(path: Path | str) -> list[Section]:
    sections = []
    for n, slide in enumerate(Presentation(str(path)).slides, start=1):
        lines = _shape_lines(slide.shapes)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                lines.append(f"Notes: {notes.text.strip()}")
        # PowerPoint stores soft line breaks as vertical tabs.
        sections.append(Section(f"slide {n}", n, "\n".join(lines).replace("\v", "\n")))
    return sections


def _shape_lines(shapes) -> list[str]:
    lines = []
    for shape in shapes:
        if isinstance(shape, GroupShape):
            lines.extend(_shape_lines(shape.shapes))
        elif shape.has_text_frame:
            if shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
        elif shape.has_table:
            # Cells covered by a merge are empty placeholders; the merge origin holds the text.
            lines.extend(_table_lines([c for c in row.cells if not c.is_spanned] for row in shape.table.rows))
    return lines


def _docx_cells(row) -> list:
    # python-docx returns a horizontally merged cell once per grid column it spans; keep one.
    cells = row.cells
    return [c for i, c in enumerate(cells) if i == 0 or c._tc is not cells[i - 1]._tc]


def _table_lines(rows: Iterable[list]) -> list[str]:
    lines = []
    for cells in rows:
        texts = [" ".join(c.text.split()) for c in cells]
        if any(texts):
            lines.append(" | ".join(texts))
    return lines


def chunk_sections(sections: list[Section], size: int, overlap: int) -> list[Chunk]:
    # Each section is split on its own so a chunk never spans two pages/slides/sections.
    splitter = RecursiveCharacterTextSplitter(chunk_size=size, chunk_overlap=overlap)
    chunks: list[Chunk] = []
    for section in sections:
        for piece in splitter.split_text(section.text):
            if piece.strip():
                chunks.append(Chunk(section.location, section.order, len(chunks), piece))
    return chunks
