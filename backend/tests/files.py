"""Builders for small test files, shared by several test modules."""

import docx
from fpdf import FPDF
from pptx import Presentation


def make_pdf(path, pages: list[str]):
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.multi_cell(0, 6, text=text)
    pdf.output(str(path))
    return path


def make_docx(path, text: str = "Some notes."):
    doc = docx.Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return path


def make_pptx(path, title: str = ""):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6] if not title else prs.slide_layouts[0])
    if title:
        slide.shapes.title.text = title
    prs.save(str(path))
    return path
