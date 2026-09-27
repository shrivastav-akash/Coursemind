import docx
from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches

from app.parsing import Section, chunk_sections, parse_docx, parse_pdf, parse_pptx


def make_pdf(path, pages: list[str]):
    pdf = FPDF()
    pdf.set_font("Helvetica", size=11)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.multi_cell(0, 6, text=text)
    pdf.output(str(path))
    return path


def test_pdf_one_section_per_page(tmp_path):
    path = make_pdf(tmp_path / "three.pdf", ["Deadlock basics.", "Paging and segmentation.", "Scheduling."])

    sections = parse_pdf(path)

    assert [(s.location, s.order) for s in sections] == [("p. 1", 1), ("p. 2", 2), ("p. 3", 3)]
    assert "Paging" in sections[1].text
    assert "\r" not in "".join(s.text for s in sections)


def test_long_page_splits_but_keeps_its_location(tmp_path):
    long_text = " ".join(f"Sentence {i} about mutual exclusion and hold and wait." for i in range(12))
    path = make_pdf(tmp_path / "long.pdf", ["Short intro.", long_text])

    chunks = chunk_sections(parse_pdf(path), size=120, overlap=20)

    page_two = [c for c in chunks if c.order == 2]
    assert len(page_two) > 1
    assert all(c.location == "p. 2" for c in page_two)
    assert all(len(c.text) <= 120 for c in chunks)
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))


def test_chunks_never_cross_sections():
    sections = [Section("p. 1", 1, "alpha " * 50), Section("p. 2", 2, "beta " * 50)]

    chunks = chunk_sections(sections, size=60, overlap=10)

    for c in chunks:
        assert ("alpha" in c.text) != ("beta" in c.text)
        assert c.location == ("p. 1" if "alpha" in c.text else "p. 2")


def test_pdf_without_text_gives_no_chunks(tmp_path):
    path = make_pdf(tmp_path / "blank.pdf", ["", ""])

    sections = parse_pdf(path)

    assert len(sections) == 2
    assert chunk_sections(sections, size=700, overlap=100) == []


def test_whitespace_only_sections_are_dropped():
    assert chunk_sections([Section("p. 1", 1, "  \n\n\t ")], size=700, overlap=100) == []


def test_docx_sections_follow_headings_and_keep_tables_in_order(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Overview before any heading.")
    doc.add_heading("Normal forms", level=1)
    doc.add_paragraph("1NF removes repeating groups.")
    table = doc.add_table(rows=3, cols=3)
    for r, row in enumerate([["Form", "Rule", "Example"], ["2NF", "No partial dependency", ""], ["", "", ""]]):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    doc.add_paragraph("After the table.")
    doc.add_heading("Transactions", level=2)
    doc.add_paragraph("ACID properties.")
    path = tmp_path / "notes.docx"
    doc.save(str(path))

    sections = parse_docx(path)

    assert [(s.location, s.order) for s in sections] == [
        ("§ (start)", 1), ("§ Normal forms", 2), ("§ Transactions", 3),
    ]
    assert sections[1].text.split("\n") == [
        "Normal forms",
        "1NF removes repeating groups.",
        "Form | Rule | Example",
        "2NF | No partial dependency | ",
        "After the table.",
    ]


def test_docx_merged_cells_and_title_without_start_section(tmp_path):
    doc = docx.Document()
    doc.add_heading("My DBMS notes", level=0)  # "Title" style
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(0, 2).text = "Right"
    path = tmp_path / "title.docx"
    doc.save(str(path))

    sections = parse_docx(path)

    assert [s.location for s in sections] == ["§ My DBMS notes"]
    assert sections[0].text.split("\n")[1] == "Merged | Right"


def test_pptx_slides_tables_nested_groups_and_notes(tmp_path):
    prs = Presentation()
    blank = prs.slide_layouts[6]

    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Deadlock"
    s1.placeholders[1].text_frame.text = "Four conditions\vfor deadlock"
    s1.notes_slide.notes_text_frame.text = "Mention Coffman."

    s2 = prs.slides.add_slide(blank)
    table = s2.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(6), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Algorithm"
    table.cell(0, 2).text = "Avoids"
    table.cell(1, 0).text = "Banker's"
    table.cell(1, 1).text = "safe state"
    table.cell(1, 2).text = "deadlock"
    inner = s2.shapes.add_group_shape().shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)).text_frame.text = "Nested group text"

    prs.slides.add_slide(blank)  # slide with no text
    path = tmp_path / "deck.pptx"
    prs.save(str(path))

    sections = parse_pptx(path)

    assert [(s.location, s.order) for s in sections] == [("slide 1", 1), ("slide 2", 2), ("slide 3", 3)]
    assert sections[0].text.split("\n") == ["Deadlock", "Four conditions", "for deadlock", "Notes: Mention Coffman."]
    assert sections[1].text.split("\n") == ["Algorithm | Avoids", "Banker's | safe state | deadlock", "Nested group text"]
    assert sections[2].text == ""
    assert {c.location for c in chunk_sections(sections, size=700, overlap=100)} == {"slide 1", "slide 2"}
